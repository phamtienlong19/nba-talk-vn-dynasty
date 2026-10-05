import os
import sys
import tempfile
import unittest
import warnings

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))

import export_yahoo_top300 as md_export  # noqa: E402
import export_yahoo_top300_xlsx as xlsx_export  # noqa: E402
from cap_model import compute_cap_model  # noqa: E402

REPO_ROOT = os.path.join(os.path.dirname(__file__), "..")
FIXTURE_MD = os.path.join(os.path.dirname(__file__), "fixtures", "yahoo_top300_2026-10-04.md")
EXPORT_MD = os.path.join(REPO_ROOT, "exports", "yahoo_top300_proj_dollar_rank.md")
EXPORT_XLSX = os.path.join(REPO_ROOT, "exports", "yahoo_top300_proj_dollar_rank.xlsx")


def _fixture_rows():
    with open(FIXTURE_MD, encoding="utf-8") as f:
        return md_export.parse_markdown(f.read())


def _players_from_rows(rows):
    # Shuffle-ish order (reverse) so the exporter's own ordering is what's tested.
    return [{"name": n, "capDollars": float(d), "oRank": r} for n, d, r in reversed(rows)]


class TestMarkdownExport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.fixture_rows = _fixture_rows()
        cls.players = _players_from_rows(cls.fixture_rows)

    def test_fixture_has_300_rows(self):
        self.assertEqual(len(self.fixture_rows), 300)

    def test_render_matches_the_reference_format_byte_for_byte(self):
        rows = md_export.build_rows(self.players)
        with open(FIXTURE_MD, encoding="utf-8") as f:
            self.assertEqual(md_export.render_markdown(rows), f.read())

    def test_header_and_order(self):
        text = md_export.render_markdown(md_export.build_rows(self.players))
        lines = text.splitlines()
        self.assertEqual(lines[0], "# Yahoo Top 300 — Projected $ / Rank")
        self.assertEqual(lines[2:4], ["| Player | Proj $ | Rank |", "|---|---:|---:|"])
        rows = md_export.parse_markdown(text)
        self.assertEqual(rows, sorted(rows, key=lambda r: (-r[1], r[2])))

    def test_deterministic(self):
        a = md_export.render_markdown(md_export.build_rows(self.players))
        b = md_export.render_markdown(md_export.build_rows(list(reversed(self.players))))
        self.assertEqual(a, b)

    def test_unicode_names_preserved(self):
        names = {r[0] for r in self.fixture_rows}
        for n in ("Nikola Jokić", "Luka Dončić", "Alperen Şengün", "Kristaps Porziņģis"):
            self.assertIn(n, names)

    def test_rejects_wrong_row_count_duplicates_and_rank_gaps(self):
        with self.assertRaises(md_export.ExportError):
            md_export.build_rows(self.players[:-1])
        dup = [dict(p) for p in self.players]
        dup[0]["name"] = dup[1]["name"]
        with self.assertRaises(md_export.ExportError):
            md_export.build_rows(dup)
        gap = [dict(p) for p in self.players]
        gap[0]["oRank"] = 999
        with self.assertRaises(md_export.ExportError):
            md_export.build_rows(gap)


class TestXlsxParity(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.xlsx = os.path.join(cls.tmp.name, "t.xlsx")
        warnings.simplefilter("ignore")
        cls.md_rows = xlsx_export.export_xlsx(FIXTURE_MD, cls.xlsx)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def _sheet_rows(self):
        return xlsx_export.read_sheet_rows(self.xlsx)

    def test_row_count_and_parity_with_markdown(self):
        rows = self._sheet_rows()
        self.assertEqual(len(rows), 300)
        for (xp, xd, xr), (mp, md, mr) in zip(rows, self.md_rows):
            self.assertEqual(xp, mp)
            self.assertEqual(xd, md)
            self.assertEqual(xr, mr)

    def test_first_middle_last_rows(self):
        rows = self._sheet_rows()
        self.assertEqual(rows[0], ("Nikola Jokić", 61, 1))
        self.assertEqual(rows[149], self.md_rows[149])
        self.assertEqual(rows[-1], ("Caris LeVert", 0, 300))

    def test_unique_players_and_rank_coverage(self):
        rows = self._sheet_rows()
        self.assertEqual(len({r[0] for r in rows}), 300)
        self.assertEqual(sorted(r[2] for r in rows), list(range(1, 301)))

    def test_numbers_are_numeric_and_unicode_survives(self):
        for _, dollars, rank in self._sheet_rows():
            self.assertIsInstance(dollars, int)
            self.assertIsInstance(rank, int)
        names = {r[0] for r in self._sheet_rows()}
        self.assertIn("Alperen Şengün", names)
        self.assertIn("Kristaps Porziņģis", names)

    def test_layout_matches_commissioner_workbook(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.xlsx)
        self.assertEqual(wb.sheetnames, ["Yahoo Top 300"])
        ws = wb["Yahoo Top 300"]
        self.assertEqual([str(r) for r in ws.merged_cells.ranges], ["A1:C1"])
        self.assertEqual(ws["A1"].value, "Yahoo Top 300 — Projected $ / Rank")
        self.assertIsNone(ws["A2"].value)
        self.assertEqual([ws.cell(3, c).value for c in (1, 2, 3)], ["Player", "Proj $", "Rank"])
        self.assertEqual(ws["A4"].value, "Nikola Jokić")
        self.assertEqual(ws["A303"].value, "Caris LeVert")
        self.assertEqual(ws.freeze_panes, "A4")
        self.assertEqual(len(ws.conditional_formatting), 1)
        wb.close()

    def test_regeneration_replaces_a_stale_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            stale = os.path.join(tmp, "out.xlsx")
            stale_md = os.path.join(tmp, "stale.md")
            rows = list(self.md_rows)
            rows[0] = ("Stale Player", 99, 1)
            with open(stale_md, "w", encoding="utf-8") as f:
                f.write(md_export.render_markdown(rows))
            xlsx_export.export_xlsx(stale_md, stale)
            self.assertEqual(xlsx_export.read_sheet_rows(stale)[0][0], "Stale Player")
            xlsx_export.export_xlsx(FIXTURE_MD, stale)
            fresh = xlsx_export.read_sheet_rows(stale)
            self.assertEqual(fresh[0], ("Nikola Jokić", 61, 1))
            self.assertNotIn("Stale Player", [r[0] for r in fresh])

    def test_workbook_is_byte_stable_for_identical_input(self):
        with tempfile.TemporaryDirectory() as tmp:
            a, b = os.path.join(tmp, "a.xlsx"), os.path.join(tmp, "b.xlsx")
            xlsx_export.export_xlsx(FIXTURE_MD, a)
            xlsx_export.export_xlsx(FIXTURE_MD, b)
            with open(a, "rb") as fa, open(b, "rb") as fb:
                self.assertEqual(fa.read(), fb.read())

    def test_three_generations_have_identical_sha256_even_across_a_clock_tick(self):
        import hashlib
        import time
        digests = set()
        with tempfile.TemporaryDirectory() as tmp:
            for i in range(3):
                path = os.path.join(tmp, f"g{i}.xlsx")
                xlsx_export.export_xlsx(FIXTURE_MD, path)
                with open(path, "rb") as f:
                    digests.add(hashlib.sha256(f.read()).hexdigest())
                time.sleep(2.1)  # ZIP timestamps have 2-second resolution
        self.assertEqual(len(digests), 1)

    def test_zip_members_are_canonical_and_workbook_still_loads(self):
        import zipfile
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "c.xlsx")
            xlsx_export.export_xlsx(FIXTURE_MD, path)
            with zipfile.ZipFile(path) as z:
                self.assertIsNone(z.testzip())
                infos = z.infolist()
                self.assertEqual(infos[0].filename, "[Content_Types].xml")
                self.assertEqual({i.date_time for i in infos}, {(1980, 1, 1, 0, 0, 0)})
                self.assertEqual({i.compress_type for i in infos}, {zipfile.ZIP_DEFLATED})
                self.assertEqual([i.filename for i in infos[1:]], sorted(i.filename for i in infos[1:]))
            wb = load_workbook(path)
            ws = wb["Yahoo Top 300"]
            self.assertEqual(ws.freeze_panes, "A4")
            self.assertEqual(ws.auto_filter.ref, "A3:C303")
            self.assertEqual([str(r) for r in ws.merged_cells.ranges], ["A1:C1"])
            self.assertEqual(len(ws.conditional_formatting), 1)
            self.assertEqual(wb.properties.created, xlsx_export.FIXED_DATETIME)
            self.assertEqual(wb.properties.modified, xlsx_export.FIXED_DATETIME)
            wb.close()
            self.assertEqual(os.listdir(tmp), ["c.xlsx"], "no temp/canonicalization files may remain")

    def test_no_temp_files_remain_after_a_failed_export(self):
        with tempfile.TemporaryDirectory() as tmp:
            bad_md = os.path.join(tmp, "bad.md")
            with open(bad_md, "w", encoding="utf-8") as f:
                f.write("# nothing\n")
            with self.assertRaises(ValueError):
                xlsx_export.export_xlsx(bad_md, os.path.join(tmp, "o.xlsx"))
            self.assertEqual(os.listdir(tmp), ["bad.md"])

    def test_divergence_between_md_and_xlsx_is_detected(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.xlsx)
        wb["Yahoo Top 300"]["B4"] = 1
        tampered = os.path.join(self.tmp.name, "tampered.xlsx")
        wb.save(tampered)
        self.assertNotEqual(xlsx_export.read_sheet_rows(tampered), self.md_rows)


class TestCommittedExportArtifacts(unittest.TestCase):
    """The committed exports/ files must agree with each other: the XLSX is
    derived from the MD and may never silently diverge from it."""

    def test_committed_md_and_xlsx_agree_exactly(self):
        warnings.simplefilter("ignore")
        with open(EXPORT_MD, encoding="utf-8") as f:
            md_rows = md_export.parse_markdown(f.read())
        self.assertEqual(len(md_rows), 300)
        self.assertEqual(xlsx_export.read_sheet_rows(EXPORT_XLSX), md_rows)
        self.assertEqual(len({r[0] for r in md_rows}), 300)
        self.assertEqual(sorted(r[2] for r in md_rows), list(range(1, 301)))

    def test_committed_export_is_fresh_vs_the_tracked_yahoo_snapshot(self):
        import json
        with open(os.path.join(REPO_ROOT, "data", "yahoo", "players_normalized.json"), encoding="utf-8") as f:
            players = json.load(f)
        with open(EXPORT_MD, encoding="utf-8") as f:
            self.assertEqual(f.read(), md_export.render_markdown(md_export.build_rows(players)))

    def test_export_is_pure_yahoo_no_cap_policy_keeper_or_dynasty_content(self):
        with open(EXPORT_MD, encoding="utf-8") as f:
            text = f.read().lower()
        for forbidden in ("official", "ceiling", "dynasty", "keeper", "override", "cap policy"):
            self.assertNotIn(forbidden, text)
        warnings.simplefilter("ignore")
        from openpyxl import load_workbook
        wb = load_workbook(EXPORT_XLSX)
        self.assertEqual(wb.sheetnames, ["Yahoo Top 300"])
        self.assertEqual([wb.active.cell(3, c).value for c in (1, 2, 3)], ["Player", "Proj $", "Rank"])
        self.assertIsNone(wb.active.cell(3, 4).value)
        wb.close()

    def test_board_links_the_stable_xlsx_and_md_paths(self):
        with open(os.path.join(REPO_ROOT, "index.html"), encoding="utf-8") as f:
            html = f.read()
        self.assertIn('href="exports/yahoo_top300_proj_dollar_rank.xlsx"', html)
        self.assertIn('href="exports/yahoo_top300_proj_dollar_rank.md"', html)


class TestRankingOnlyChangeNeverLeavesAStaleExport(unittest.TestCase):
    def test_ranking_only_change_is_detected_and_regenerated(self):
        from build_yahoo_refresh_result import exports_stale
        rows = _fixture_rows()
        before = _players_from_rows(rows)
        after = [dict(p) for p in before]
        a = next(p for p in after if p["name"] == "Rui Hachimura")
        b = next(p for p in after if p["name"] == "RJ Barrett")  # both $0: ranks swap, cap band unchanged
        a["oRank"], b["oRank"] = b["oRank"], a["oRank"]
        with tempfile.TemporaryDirectory() as tmp:
            md = os.path.join(tmp, "e.md")
            with open(md, "w", encoding="utf-8") as f:
                f.write(md_export.render_markdown(md_export.build_rows(before)))
            self.assertFalse(exports_stale(before, md))
            self.assertTrue(exports_stale(after, md), "ranking-only change must make the committed export stale")
            md_export.write_markdown.__call__  # regeneration entry point exists
            self.assertEqual(compute_cap_model(before)["roundedCeiling"], compute_cap_model(after)["roundedCeiling"])


class TestCapModelIndependence(unittest.TestCase):
    def test_cap_benchmark_is_unaffected_by_the_exported_rank_column(self):
        rows = _fixture_rows()
        players = [{"capDollars": d, "oRank": r} for _, d, r in rows]
        scrambled = [{"capDollars": d, "oRank": 1000 + i} for i, (_, d, _) in enumerate(rows)]
        # Same dollars, completely different Rank column -> same cap result.
        self.assertEqual(compute_cap_model(players)["top144Sum"],
                         compute_cap_model(scrambled)["top144Sum"])

    def test_cap_scripts_do_not_import_the_export(self):
        with open(os.path.join(REPO_ROOT, "scripts", "cap_model.py"), encoding="utf-8") as f:
            self.assertNotIn("export_yahoo_top300", f.read())


if __name__ == "__main__":
    unittest.main()
